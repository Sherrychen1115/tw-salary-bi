# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "openpyxl>=3.1",
# ]
# ///
"""驗證 data/salary_long.csv 是否忠實反映原始 xlsx。

用法：
    uv run scripts/validate.py              # 每次抽不同的 20 格
    uv run scripts/validate.py --seed 42    # 固定亂數種子，可重現

檢查項目：
  1. 列數 = 指標數 × 年數 × 行業數，且 (指標, 年度, 行業) 不重複
  2. 隨機抽 20 格，與原始 xlsx 逐格比對（另核對有值／無資料的總格數）
  3. 「受僱員工人數」：各層子行業加總 ≈ 上層（中類→大類、大類→部門、小類→中類）

這支程式自己重新讀 xlsx，不引用 etl_salary.py 的任何邏輯。
每次執行都會把結果寫到 data/validation_report.md（每項檢查的預期值、實際值、是否通過）。
全部通過時結束碼為 0，有任何不符為 1。
"""

from __future__ import annotations

import argparse
import csv
import random
import sys
import warnings
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SRC_DIR = ROOT / "行政院主計總處統計資料"
CSV_PATH = ROOT / "data" / "salary_long.csv"
REPORT_PATH = ROOT / "data" / "validation_report.md"

FIELDS = ["indicator", "unit", "year_roc", "year_ad", "industry", "level", "parent", "value"]
SAMPLE_SIZE = 20
EMPLOYEE_INDICATOR = "受僱員工人數"

results: list[dict] = []  # 每項檢查一筆，最後寫進報告
samples: list[dict] = []  # 抽樣比對的每一格
section = ""


def begin(title: str) -> None:
    global section
    section = title
    print(f"\n{title}")


def check(name: str, expected, actual, ok: bool | None = None) -> None:
    """記錄一項檢查；沒給 ok 時以 expected == actual 判定。"""
    if ok is None:
        ok = expected == actual
    print(f"  [{'通過' if ok else '不符'}] {name}：預期 {expected}，實際 {actual}")
    results.append(
        {"section": section, "name": name, "expected": expected, "actual": actual, "ok": ok}
    )


def md(text) -> str:
    return str(text).replace("|", "\\|")


def write_report(src: Path, seed: int) -> None:
    failed = [r for r in results if not r["ok"]]
    verdict = "全部通過" if not failed else f"{len(failed)} 項不符"
    lines = [
        "# 驗證報告",
        "",
        "由 `uv run scripts/validate.py` 產生，請勿手動編輯。",
        "",
        f"- 執行時間：{datetime.now():%Y-%m-%d %H:%M:%S}",
        f"- 原始檔：`{src.parent.name}/{src.name}`",
        f"- 受檢檔：`{CSV_PATH.relative_to(ROOT).as_posix()}`",
        f"- 抽樣亂數種子：{seed}（用 `--seed {seed}` 可重現同一批抽樣）",
        f"- 結果：**{verdict}**（{len(results) - len(failed)} / {len(results)} 項通過）",
    ]
    current = None
    for r in results:
        if r["section"] != current:
            current = r["section"]
            lines += [
                "",
                f"## {current}",
                "",
                "| 檢查項目 | 預期值 | 實際值 | 結果 |",
                "|---|---|---|---|",
            ]
        lines.append(
            f"| {md(r['name'])} | {md(r['expected'])} | {md(r['actual'])} "
            f"| {'通過' if r['ok'] else '**不符**'} |"
        )
    lines += [
        "",
        f"## 抽樣明細（{len(samples)} 格）",
        "",
        "預期值是原始 xlsx 的儲存格內容（「-」代表無資料），實際值是 CSV 的 value（無資料時為空）。",
        "",
        "| 儲存格 | 指標 | 年度 | 行業 | 預期值（xlsx） | 實際值（CSV） | 結果 |",
        "|---|---|---|---|---|---|---|",
    ]
    for s in samples:
        lines.append(
            f"| {s['coord']} | {md(s['indicator'])} | {s['year']} | {md(s['industry'])} "
            f"| {md(s['xlsx'])} | {md(s['csv'])} | {'通過' if s['ok'] else '**不符**'} |"
        )
    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")


def load_csv() -> list[dict]:
    raw = CSV_PATH.read_bytes()
    check("CSV 開頭 3 個位元組（UTF-8 BOM）", "EF BB BF", raw[:3].hex(" ").upper())
    with CSV_PATH.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        check("欄位名稱與順序", ", ".join(FIELDS), ", ".join(reader.fieldnames or []))
        return list(reader)


def load_sheet(path: Path):
    """讀第一個工作表，回傳 (cells, 資料列號, 資料欄號)；cells[(row, col)] 為原始儲存格值。"""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb.worksheets[0]
        cells = {}
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            for c, v in enumerate(row, start=1):
                if v not in (None, ""):
                    cells[(r, c)] = v
    finally:
        wb.close()

    max_row = max(r for r, _ in cells)
    note_row = next(
        (r for r in range(4, max_row + 1) if str(cells.get((r, 1), "")).strip().startswith("註")),
        max_row + 1,
    )
    data_rows = [r for r in range(4, note_row) if (r, 2) in cells]
    data_cols = sorted(c for (r, c) in cells if r == 1 and c >= 3)
    return cells, data_rows, data_cols


def indicator_of(cells, row: int) -> str:
    """A 欄是合併儲存格，往上找到第一個有值的儲存格。"""
    while (row, 1) not in cells:
        row -= 1
    return str(cells[(row, 1)]).strip()


def same_value(xlsx_value, csv_value: str) -> bool:
    if xlsx_value is None or xlsx_value == "-":
        return csv_value == ""
    if csv_value == "":
        return False
    return float(csv_value) == float(xlsx_value)


def raw_indicator(r: dict) -> str:
    """以「指標(單位)」還原出 xlsx 上的原始指標名稱，當作比對用的 key。"""
    return f"{r['indicator']}({r['unit']})" if r["unit"] else r["indicator"]


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx", nargs="?", help="原始 xlsx（預設讀資料夾內檔名排序最後一個）")
    parser.add_argument("--seed", type=int, default=None, help="抽樣用的亂數種子")
    args = parser.parse_args()

    if args.xlsx:
        src = Path(args.xlsx)
    else:
        src = sorted(p for p in SRC_DIR.glob("*.xlsx") if not p.name.startswith("~$"))[-1]
    seed = args.seed if args.seed is not None else random.randrange(1_000_000)
    print(f"原始檔：{src.name}")
    print(f"輸出檔：{CSV_PATH.relative_to(ROOT)}")

    begin("0. 檔案格式")
    rows = load_csv()
    cells, data_rows, data_cols = load_sheet(src)

    index: dict[tuple, dict] = {}
    duplicates = 0
    for r in rows:
        key = (raw_indicator(r), r["year_roc"], r["industry"])
        duplicates += key in index
        index[key] = r

    begin("1. 列數")
    row_indicator = {r: indicator_of(cells, r) for r in data_rows}
    xlsx_indicators = set(row_indicator.values())
    csv_indicators = {raw_indicator(r) for r in rows}
    n_ind = len(xlsx_indicators)
    n_year = len({cells[(r, 2)] for r in data_rows})
    n_industry = len(data_cols)
    check(
        f"列數 = 指標 {n_ind} × 年度 {n_year} × 行業 {n_industry}",
        n_ind * n_year * n_industry,
        len(rows),
    )
    check("（指標, 年度, 行業）重複筆數", 0, duplicates)
    check("指標數", n_ind, len(csv_indicators))
    check("年度數", n_year, len({r["year_roc"] for r in rows}))
    check("行業數", n_industry, len({r["industry"] for r in rows}))
    check(
        "year_ad ≠ year_roc + 1911 的列數",
        0,
        sum(int(r["year_ad"]) != int(r["year_roc"]) + 1911 for r in rows),
    )
    check(
        "指標名稱不在 xlsx 內的列數（例如誤讀「註:」之後的列）",
        0,
        sum(raw_indicator(r) not in xlsx_indicators for r in rows),
    )

    begin(f"2. 隨機抽 {SAMPLE_SIZE} 格逐格比對（seed={seed}）")
    rng = random.Random(seed)
    population = [(r, c) for r in data_rows for c in data_cols]
    mismatches = 0
    for r, c in rng.sample(population, SAMPLE_SIZE):
        ind = row_indicator[r]
        year = str(cells[(r, 2)]).strip().removesuffix("年")
        industry = str(cells[(1, c)]).strip()
        xlsx_value = cells.get((r, c))
        rec = index.get((ind, year, industry))
        csv_value = None if rec is None else rec["value"]
        ok = rec is not None and same_value(xlsx_value, csv_value)
        mismatches += not ok
        coord = f"{get_column_letter(c)}{r}"
        samples.append(
            {
                "coord": coord,
                "indicator": ind,
                "year": f"{year}年",
                "industry": industry,
                "xlsx": xlsx_value,
                "csv": "找不到這一列" if rec is None else (csv_value or "（空）"),
                "ok": ok,
            }
        )
        print(
            f"    {'✓' if ok else '✗'} {coord:<6} {ind}｜{year}年｜{industry}"
            f"：xlsx={xlsx_value!r} csv={csv_value!r}"
        )
    check("抽樣相符格數", SAMPLE_SIZE, SAMPLE_SIZE - mismatches)

    xlsx_numeric = sum(1 for p in population if isinstance(cells.get(p), (int, float)))
    csv_numeric = sum(1 for r in rows if r["value"] != "")
    check("有值的格數（全表）", xlsx_numeric, csv_numeric)
    check("無資料的格數（全表）", len(population) - xlsx_numeric, len(rows) - csv_numeric)
    check(
        "value 為 0 的格數（「-」不可寫成 0）",
        sum(1 for p in population if cells.get(p) == 0),
        sum(1 for r in rows if r["value"] != "" and float(r["value"]) == 0),
    )
    xlsx_has_value = {
        row_indicator[r] for r, c in population if isinstance(cells.get((r, c)), (int, float))
    }
    csv_has_value = {raw_indicator(r) for r in rows if r["value"] != ""}
    check(
        "整個指標都無資料的指標",
        "、".join(sorted(xlsx_indicators - xlsx_has_value)) or "無",
        "、".join(sorted(csv_indicators - csv_has_value)) or "無",
    )

    begin(f"3. {EMPLOYEE_INDICATOR}：子行業加總 ≈ 上層行業")
    emp = [r for r in rows if r["indicator"] == EMPLOYEE_INDICATOR]
    check(f"指標「{EMPLOYEE_INDICATOR}」的列數", n_year * n_industry, len(emp))
    value = {(r["industry"], r["year_roc"]): r["value"] for r in emp}
    level = {r["industry"]: r["level"] for r in emp}
    children = defaultdict(set)
    for r in emp:
        if r["parent"]:
            children[r["parent"]].add(r["industry"])
    check("parent 不是 CSV 內行業的個數", 0, sum(p not in level for p in children))
    years = sorted({r["year_roc"] for r in emp})
    by_level = defaultdict(lambda: [0, 0, 0])  # 上層 level -> [比對組數, 不符組數, 最大差]
    for parent, kids in sorted(children.items()):
        # 原始數字是四捨五入後的整數，每個子行業最多差 0.5，容許差 = 子行業數
        tolerance = len(kids)
        for y in years:
            total = value.get((parent, y), "")
            parts = [value[(k, y)] for k in kids]
            if total == "" or "" in parts:
                continue
            diff = abs(float(total) - sum(float(p) for p in parts))
            stat = by_level[level.get(parent, "?")]
            stat[0] += 1
            stat[2] = max(stat[2], diff)
            if diff > tolerance:
                stat[1] += 1
                print(f"    ✗ {parent} {y}年：上層 {total}，子行業加總差 {diff:g}")
    for parent_lv, child_lv in (("大類", "中類"), ("部門", "大類"), ("中類", "小類")):
        n, bad, max_diff = by_level[parent_lv]
        check(
            f"{child_lv}加總 ≈ {parent_lv}（每個上層行業 × 年度為一組）",
            "至少 1 組可比對，且每組的差 ≤ 子行業數",
            f"比對 {n} 組，不符 {bad} 組，最大差 {max_diff:g} 人",
            ok=n > 0 and bad == 0,
        )

    write_report(src, seed)
    failed = [r for r in results if not r["ok"]]
    print(f"\n報告：{REPORT_PATH.relative_to(ROOT)}")
    if failed:
        print(f"驗證失敗：{len(failed)} 項不符")
        return 1
    print("驗證全部通過")
    return 0


if __name__ == "__main__":
    sys.exit(main())
